# Builds pink-sheet.sample.xlsx: a SMALL MADE-UP TEST FIXTURE shaped like the World Bank
# "CMO-Historical-Data-Monthly.xlsx" workbook (sheet "Monthly Prices": title rows, a names row,
# a units row, then one row per month "YYYYMmm"). The numbers are invented for tests.
# They are NOT real prices and must never be published.
# Run: python3 make-pink-sheet.py   (needs openpyxl; output is committed, so tests don't need Python)
import os
from openpyxl import Workbook

cols = [
    ("Crude oil, average", "($/bbl)", 70.0),
    ("Coal, Australian", "($/mt)", 110.0),
    ("Coal, South African **", "($/mt)", 95.0),
    ("Natural gas, Europe", "($/mmbtu)", 11.0),
    ("Liquefied natural gas, Japan", "($/mmbtu)", 12.0),
    ("Cocoa", "($/kg)", 6.5),
    ("Coffee, Arabica", "($/kg)", 7.4),
    ("Coffee, Robusta", "($/kg)", 4.6),
    ("Palm oil", "($/mt)", 980.0),
    ("Soybeans", "($/mt)", 420.0),
    ("Maize", "($/mt)", 200.0),
    ("Rice, Thai 5% ", "($/mt)", 400.0),
    ("Wheat, US SRW", "($/mt)", 230.0),
    ("Wheat, US HRW", "($/mt)", 250.0),
    ("Sugar, world", "($/kg)", 0.40),
    ("Cotton, A Index", "($/kg)", 1.70),
    ("Phosphate rock", "($/mt)", 150.0),
    ("DAP", "($/mt)", 700.0),
    ("TSP", "($/mt)", 600.0),
    ("Urea ", "($/mt)", 420.0),
    ("Potassium chloride **", "($/mt)", 350.0),
    ("Aluminum", "($/mt)", 2600.0),
    ("Iron ore, cfr spot", "($/dmtu)", 100.0),
    ("Copper", "($/mt)", 10000.0),
    ("Lead", "($/mt)", 2000.0),
    ("Tin", "($/mt)", 33000.0),
    ("Nickel", "($/mt)", 15500.0),
    ("Zinc", "($/mt)", 2800.0),
    ("Gold", "($/troy oz)", 3500.0),
    ("Platinum", "($/troy oz)", 1300.0),
    ("Silver", "($/troy oz)", 40.0),
]

wb = Workbook()
wb.active.title = "Index"
wb.active["A1"] = "SAMPLE DATA - NOT REAL PRICES - TEST FIXTURE"
ws = wb.create_sheet("Monthly Prices")
ws["A1"] = "SAMPLE DATA - NOT REAL PRICES - TEST FIXTURE (shaped like World Bank Commodity Price Data)"
ws["A2"] = "Monthly prices in nominal US dollars, 1960 to present"
ws["A3"] = "(monthly series are available only in nominal US dollar terms)"
for j, (name, unit, _) in enumerate(cols):
    ws.cell(row=5, column=j + 2, value=name)
    ws.cell(row=6, column=j + 2, value=unit)
months = [f"2025M{m:02d}" for m in range(1, 13)] + [f"2026M{m:02d}" for m in range(1, 10)]
for i, mo in enumerate(months):
    r = 7 + i
    ws.cell(row=r, column=1, value=mo)
    for j, (name, unit, base) in enumerate(cols):
        # Deterministic made-up wiggle; one deliberate gap ("…") to test missing values.
        if name == "Tin" and mo == "2025M03":
            ws.cell(row=r, column=j + 2, value="…")
            continue
        f = 1 + 0.01 * (((i * 7 + j * 3) % 11) - 5)
        ws.cell(row=r, column=j + 2, value=round(base * f, 4))
wb.save(os.path.join(os.path.dirname(os.path.abspath(__file__)), "pink-sheet.sample.xlsx"))
print("wrote pink-sheet.sample.xlsx")
